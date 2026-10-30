"""Build brownfield-modernization.pptx from the HTML deck content, with Vietnamese speaker notes.

Run: uv run --no-project --with python-pptx python docs/presentations/build_brownfield_pptx.py
Slide copy lives in this script; keep it in step with brownfield-modernization.html.
The slide 4 lifecycle diagram is read from the HTML and redrawn as native shapes.
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

C = {k: RGBColor.from_string(v) for k, v in {
    "bg": "EEF1F4", "surface": "FFFFFF", "ink": "141C26", "muted": "4B5563", "line": "CFD6DF",
    "agent": "2563EB", "human": "9A4508", "human_soft": "FBEADB", "success": "15803D",
    "success_soft": "DCFCE7", "band_people": "F7F1EA", "band_agents": "EAF0FA", "band_auto": "E9ECF0", "band_col": "F6F7F9",
    "agent_border": "B3C8F5", "human_border": "DDB89C", "success_border": "A9D4B9",
}.items()}
DISPLAY, BODY, MONO = "Aptos Display", "Aptos", "Consolas"

W, H = 13.333, 7.5
L, CW = 0.6, 13.333 - 1.2
GAP = 0.22

HERE = Path(__file__).resolve().parent
DECK_HTML = HERE / "brownfield-modernization.html"
DECK_PPTX = HERE / "brownfield-modernization.pptx"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def runs(paragraph, text, size, color, font=BODY, bold=False):
    """Add runs to a paragraph; **bold** and `code` spans are styled inline."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        r = paragraph.add_run()
        is_bold, is_code = part.startswith("**"), part.startswith("`")
        r.text = part[2:-2] if is_bold else part[1:-1] if is_code else part
        r.font.size = Pt(size)
        r.font.name = MONO if is_code else font
        r.font.bold = bold or is_bold
        r.font.color.rgb = color


def bullet(paragraph, kind="char", indent=0):
    pPr = paragraph._p.get_or_add_pPr()
    mar = 0.24 + indent * 0.26
    pPr.set("marL", str(int(Inches(mar))))
    pPr.set("indent", str(int(-Inches(0.2))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if kind == "num":
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(pPr, qn("a:buChar")).set("char", "–" if indent else "•")


def spacing(paragraph, before=0, line=1.15):
    paragraph.space_before = Pt(before)
    paragraph.line_spacing = line


def frame(shape, margin=0.18, anchor=MSO_ANCHOR.TOP, wrap=True):
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin * 0.8)
    return tf


def box(slide, x, y, w, h, fill, border, weight=1.0, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if border is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = border
        s.line.width = Pt(weight)
    return s


def text(slide, x, y, w, h, lines, size, color, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = frame(tb, 0, anchor)
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs(p, line, size, color, font, bold)
    return tb


def card(slide, x, y, w, h, title, items=(), paras=(), kind=None, fill=None):
    """A card: title, optional paragraphs, then bullet items (str, or (str, indent))."""
    s = box(slide, x, y, w, h, fill or C["bg"], C["line"])
    tf = frame(s, 0.22)
    p = tf.paragraphs[0]
    runs(p, title, 20, C["ink"], DISPLAY, True)
    spacing(p)
    for para in paras:
        p = tf.add_paragraph()
        runs(p, para, 16, C["ink"])
        spacing(p, 8, 1.2)
    for item in items:
        indent = 0
        if isinstance(item, tuple):
            item, indent = item
        p = tf.add_paragraph()
        runs(p, item, 16, C["ink"])
        bullet(p, kind or "char", indent)
        spacing(p, 8 if not indent else 3, 1.2)
    return s


def callout(slide, x, y, w, h, label, body, tone="agent"):
    fill, border, accent = {
        "agent": (C["band_agents"], C["agent_border"], C["agent"]),
        "human": (C["human_soft"], C["human_border"], C["human"]),
        "success": (C["success_soft"], C["success_border"], C["success"]),
    }[tone]
    s = box(slide, x, y, w, h, fill, border)
    tf = frame(s, 0.2, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    runs(p, label + " ", 17, accent, BODY, True)
    runs(p, body, 17, C["ink"])
    return s


def header(slide, kicker, title, subtitle):
    text(slide, L, 0.38, CW, 0.3, kicker.upper(), 12, C["agent"], MONO, True)
    text(slide, L, 0.68, CW, 0.62, title, 30, C["ink"], DISPLAY, True, anchor=MSO_ANCHOR.BOTTOM)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L), Inches(1.4), Inches(L + CW), Inches(1.4))
    ln.line.color.rgb = C["line"]
    ln.line.width = Pt(1)
    text(slide, L, 1.5, CW, 0.4, subtitle, 18, C["muted"], DISPLAY)


def cols(n, y, h):
    w = (CW - GAP * (n - 1)) / n
    return [(L + i * (w + GAP), y, w, h) for i in range(n)]


def table(slide, rows, y, widths, head_h, row_h):
    tbl = slide.shapes.add_table(len(rows), len(widths), Inches(L), Inches(y), Inches(CW),
                                 Inches(head_h + row_h * (len(rows) - 1))).table
    for i, wdt in enumerate(widths):
        tbl.columns[i].width = Inches(wdt)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(head_h if r == 0 else row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C["bg"] if r == 0 else C["surface"]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.15)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            runs(tf.paragraphs[0], val, 16, C["ink"], DISPLAY if r == 0 else BODY, r == 0 or c == 0)
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    etree.SubElement(tblPr, qn("a:tableStyleId")).text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # Table Grid
    return tbl


SVG_FILL = {
    "band-people": ("band_people", None), "band-agents": ("band_agents", None), "band-auto": ("band_auto", None),
    "strip-people": ("human", None), "strip-agents": ("agent", None), "strip-auto": ("muted", None),
    "n-agent": ("surface", "agent"), "n-opt": ("surface", "agent"), "n-human": ("human_soft", "human"),
    "n-human-opt": ("human_soft", "human"), "n-hook": ("surface", "muted"), "chip": ("ink", None),
    "chip-out": ("surface", "ink"), "gate-band": ("bg", None), "g-human": ("human", None),
    "g-agent": ("agent", None), "g-derived": ("surface", "agent"),
}
SVG_TEXT = {  # class: (size units, color, font, bold)
    "stage": (15, "ink", DISPLAY, True), "lane": (13, "ink", DISPLAY, True), "lane-sub": (11, "muted", BODY, False),
    "t": (12, "ink", BODY, False), "t-mono": (12, "ink", MONO, False), "lbl": (11, "muted", BODY, False),
    "chip-t": (12, "bg", MONO, False), "gate": (12, "ink", MONO, True),
}
SVG_LINE = {"sep": ("line", None), "col": ("line", MSO_LINE.SQUARE_DOT), "a": ("muted", None),
            "a-dot": ("muted", MSO_LINE.ROUND_DOT), "a-wait": ("human", None)}


def svg_diagram(slide, svg, x0, y0, max_w, max_h, font_k=1.0):
    """Redraw a class-styled SVG diagram as native, editable shapes."""
    root = etree.fromstring(svg.encode())
    vb = [float(v) for v in root.get("viewBox").split()]
    k = min(max_w / vb[2], max_h / vb[3])
    ox = x0 + (max_w - vb[2] * k) / 2
    num = lambda el, a: float(el.get(a, 0))
    for el in root.iter(etree.Element):
        tag, cls = etree.QName(el).localname, el.get("class", "")
        if tag == "rect" and cls in SVG_FILL:
            fill, stroke = SVG_FILL[cls]
            x, y, w, h = (num(el, a) for a in ("x", "y", "width", "height"))
            if "rotate" in (el.get("transform") or ""):
                d = w * 1.414
                x, y, w, h = x + w / 2 - d / 2, y + h / 2 - d / 2, d, d
                shape = MSO_SHAPE.DIAMOND
            else:
                shape = MSO_SHAPE.ROUNDED_RECTANGLE if el.get("rx") else MSO_SHAPE.RECTANGLE
            sh = box(slide, ox + x * k, y0 + y * k, w * k, h * k, C[fill], C[stroke] if stroke else None, 1.25,
                     min(0.5, num(el, "rx") / max(min(w, h), 1)), shape)
            if cls.endswith("-opt"):
                sh.line.dash_style = MSO_LINE.DASH
        elif tag in ("line", "polyline") and cls in SVG_LINE:
            color, dash = SVG_LINE[cls]
            if tag == "line":
                pts = [(num(el, "x1"), num(el, "y1")), (num(el, "x2"), num(el, "y2"))]
            else:
                vals = [float(v) for v in re.split(r"[ ,]+", el.get("points").strip())]
                pts = list(zip(vals[::2], vals[1::2]))
            for i, ((xa, ya), (xb, yb)) in enumerate(zip(pts, pts[1:])):
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ox + xa * k), Inches(y0 + ya * k),
                                                Inches(ox + xb * k), Inches(y0 + yb * k))
                ln.line.color.rgb = C[color]
                ln.line.width = Pt(1.5 if cls == "a-wait" else 1)
                if dash:
                    ln.line.dash_style = dash
                if el.get("marker-end") and i == len(pts) - 2:
                    etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
        elif tag == "text" and cls in SVG_TEXT:
            size_u, color, font, bold = SVG_TEXT[cls]
            size = size_u * k * 72 * font_k
            content = "".join(el.itertext())
            w = len(content) * size * 0.62 / 72 + 0.1
            anchor = el.get("text-anchor", "start")
            x = ox + num(el, "x") * k
            left = x - w / 2 if anchor == "middle" else x - w if anchor == "end" else x
            align = {"middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}.get(anchor, PP_ALIGN.LEFT)
            top = y0 + num(el, "y") * k - size * 1.02 / 72
            tb = text(slide, left, top, w, size * 1.3 / 72, content, size, C[color], font, bold, align)
            tb.text_frame.word_wrap = False


def new_slide(notes):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["surface"]
    s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def diagram_scale(units_w):
    return CW / units_w


def dbox(slide, ox, oy, k, x, y, w, h, fill, border, weight, lines, font_k=1.05):
    """Diagram box from SVG units; lines are (text, size_units, color, font, bold)."""
    s = box(slide, ox + x * k, oy + y * k, w * k, h * k, fill, border, weight, 0.12)
    tf = frame(s, 0.02, MSO_ANCHOR.MIDDLE)
    for i, (t, size, color, font, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        runs(p, t, size * k * 72 * font_k, color, font, bold)
        spacing(p, 0, 1.0)
    return s


# ---------- Slide 1 ----------
s = new_slide("""
Mở đầu: đây là hướng dẫn vận hành end-to-end để hiện đại hoá hệ thống brownfield bằng gin-workflow và SDD (Spec-Driven Development), từ khâu khám phá hệ thống legacy đến giao hàng tăng dần.

Ba trụ cột:
- Evidence Discovery: phân tích codebase tự động, thay cho việc dựa vào "kiến thức truyền miệng".
- Contract-First SDD: spec (spec.md) là hợp đồng sống, được xác nhận qua pull request trước khi viết code.
- Gated Execution: mỗi track làm trong Git worktree riêng, làm theo TDD, có review độc lập ghi vào review ledger và audit trace tự động trước khi ship.
""")
text(s, L, 1.55, CW, 0.3, "GIN-WORKFLOW 2.1 — MODERNIZATION ARCHITECTURE", 13, C["agent"], MONO, True)
text(s, L, 1.95, CW * 0.85, 1.5, "Modernizing Brownfield Systems with gin-workflow & SDD", 42, C["ink"], DISPLAY, True)
text(s, L, 3.5, CW, 0.45, "From legacy discovery to incremental, verified delivery.", 20, C["muted"], DISPLAY)
for (x, y, w, h), (t, b) in zip(cols(3, 4.35, 1.6), [
    ("1. Evidence Discovery", "Baseline the legacy system from code, not memory."),
    ("2. Contract-First SDD", "Specs confirmed by pull request before any code is written."),
    ("3. Gated Execution", "Isolated worktrees, TDD, independent review, and trace audits."),
]):
    card(s, x, y, w, h, t, paras=[b])

# ---------- Slide 2 ----------
s = new_slide("""
Vấn đề của hệ thống legacy là "khoảng trống tri thức":
- Người viết ban đầu đã rời đi; business rule chỉ còn nằm trong code.
- Wiki, README, sơ đồ chậm hơn production nhiều năm.
- Có nhiều thứ ẩn: database trigger ngầm, endpoint không có tài liệu, các lời gọi chằng chịt.

/tech-doc giải quyết bằng bằng chứng:
- Quét trực tiếp cây code, database schema và API handler.
- Xuất ra 10 module có cấu trúc trong docs/codebase/, ví dụ 03-ARCHITECTURE.md, 05-API.md, 07-STACK.md, 04-STRUCTURE.md.
- Mọi kết luận kiến trúc đều trích dẫn đúng file và dòng, không suy đoán.

Kết quả: một baseline có thể kiểm chứng, làm trong vài giờ thay vì vài tuần, dùng làm đầu vào cho roadmap.
""")
header(s, "Phase 0: Discovery", "Discover the Legacy System from Evidence (`/tech-doc`)",
       "Architecture and dependency baselines taken directly from the repository.")
a, b = cols(2, 2.1, 3.25)
card(s, *a, "The Knowledge Gap", items=[
    "Original authors are gone; rules live only in code.",
    "Docs and diagrams lag years behind production.",
    "Hidden triggers, endpoints, and coupling."])
card(s, *b, "What `/tech-doc` Produces", items=[
    "Scans code, schemas, and API handlers.",
    "10 baseline modules in `docs/codebase/`.",
    "Every conclusion cites a file and line range."])
callout(s, L, 5.6, CW, 0.75, "Outcome:", "A verifiable baseline of current behavior that seeds the roadmap.")

# ---------- Slide 3 ----------
s = new_slide("""
Sắp xếp roadmap theo phụ thuộc:
- Output của /tech-doc cho thấy call graph và phụ thuộc dữ liệu giữa các module legacy.
- Ưu tiên các module không có hoặc có rất ít phụ thuộc upstream, để không bị kẹt về kiến trúc.
- Đầu tư vào spec là xứng đáng: code trên spec chưa xác nhận sẽ gây rework lớn về sau.

Hai luồng chạy song song, không ai phải ngồi chờ:
- Engineering: chạy ổn định app legacy trong Docker, dựng repo backend/frontend mới, thiết lập CI/CD và reverse proxy.
- BA: tập trung vào business logic, edge case và acceptance scenario cho module đầu tiên không bị chặn.

Mốc theo sprint:
- Sprint 1: môi trường dev chạy được và một spec đã xác nhận, sẵn sàng triển khai.
- Sprint 2: vertical slice hiện đại đầu tiên được ship sau facade; spec tiếp theo đã sẵn sàng.
- Sprint 3 trở đi: giao hàng liên tục, có regression test bảo vệ, dần loại bỏ các đường code legacy đã được thay thế (strangler pattern).

Nhấn mạnh quy tắc: không bắt đầu code khi spec chưa được xác nhận.
""")
header(s, "Roadmap Governance", "Sequence by Dependency, Start in Parallel",
       "Unblocked modules go first; engineering setup and spec work run side by side.")
rows = [
    ("Sprint", "Engineering Stream", "Product & Analysis Stream"),
    ("Sprint 1", "Run legacy in Docker; scaffold new repos; set up CI/CD & proxy.", "Pick the first unblocked module; confirm its spec via PR."),
    ("Sprint 2", "Plan tracks, open worktrees, build the first epic with TDD.", "Spec the next dependent epic."),
    ("Sprint 3+", "Build the next epics; retire replaced legacy paths.", "Keep living specs current via `spec-delta.md`."),
]
table(s, rows, 2.1, (1.9, (CW - 1.9) / 2, (CW - 1.9) / 2), 0.6, 0.83)
callout(s, L, 5.55, CW, 0.75, "Rule:", "No code starts on an unconfirmed spec.")

# ---------- Slide 4 ----------
s = new_slide("""
Sơ đồ swim-lane cho thấy ai làm gì ở từng stage: discuss, plan, orchestrate, execute (theo từng track) và ship. Hàng dưới cùng là gate đóng mỗi stage.

Các làn:
- Con người: Requester (BA/PM), Area lead (sở hữu vùng code), Developer (chạy các stage), Ship approver (trong team mode, ví dụ QE).
- Agent: Main agent (phiên làm việc của bạn), Worker (gin-orchestrator, tuỳ chọn), Reviewer (phiên hoặc provider độc lập).
- Tự động hoá: git hooks và CI, mức kiểm tra tuỳ theo rigor.

Đi theo từng stage:
- discuss: agent hỏi, đề xuất và viết spec; BA/PM xác nhận spec → gate requirement_confirmed.
- plan: agent viết plan gồm các track, vai trò và mức reasoning; area lead duyệt → gate plan_approved.
- orchestrate: agent tạo epic, track bead, route và worktree → gate orchestration_ready.
- execute: với mỗi track, agent nhận track và viết test trước; worker có thể triển khai track được route; reviewer ghi review ledger với số vòng giới hạn; agent sửa findings rồi đóng track. Track cần con người thì developer làm. pre-commit chạy lint và typecheck; pre-push chạy test, build, e2e. Gate implementation_complete tự suy ra khi mọi track đã đóng.
- ship: developer chọn merge, mở PR, giữ lại hoặc huỷ; ship approver duyệt PR; agent kiểm tra ledger, merge và dọn dẹp → gate shipped.

Đường nét cam là chỗ agent phải dừng chờ con người.
""")
text(s, L, 0.38, CW, 0.3, "TEAM GOVERNANCE", 12, C["agent"], MONO, True)
text(s, L, 0.68, CW, 0.62, "Who Does What, Stage by Stage", 30, C["ink"], DISPLAY, True, anchor=MSO_ANCHOR.BOTTOM)
svg_diagram(s, re.search(r'<svg class="lifecycle".*?</svg>', DECK_HTML.read_text(encoding="utf-8"), re.S).group(0), L, 1.5, CW, 5.85, 1.3)

# ---------- Slide 5 ----------
s = new_slide("""
Spec do AI soạn được kiểm chứng qua ba lớp:

1. Ngữ pháp chặt chẽ:
- Cú pháp EARS: mỗi requirement phải nêu rõ điều kiện bằng SHALL hoặc MUST.
- Mỗi requirement cần ít nhất một scenario GIVEN / WHEN / THEN cụ thể.
- specs lint báo lỗi ngay nếu thiếu cú pháp hoặc thiếu scenario.

2. Cổng PR do con người duyệt:
- Thay đổi spec được đề xuất trên nhánh spec/<epic> và gửi thành pull request.
- Gate requirement-confirmed yêu cầu BA duyệt trên GitHub/GitLab.
- Chưa merge spec PR thì chưa được lập kế hoạch hay viết code.

3. Bám vào giao diện thật:
- BA lưu ảnh chụp màn hình vào docs/specs/<cap>/assets/.
- Agent đọc wireframe để rút ra các trường nhập liệu, quy tắc validate và trạng thái lỗi, rồi liên kết ảnh vào spec.

Thông điệp chính: agent soạn và cấu trúc requirement, còn BA là người duy nhất có quyền xác nhận.
""")
header(s, "Stage 1: Discuss", "How AI-Drafted Specs Get Verified",
       "Strict grammar, grounded inputs, and a human approval gate.")
for (x, y, w, h), (t, ps) in zip(cols(3, 2.1, 3.25), [
    ("1. Strict Grammar", ["EARS `SHALL`/`MUST` statements, each with a GIVEN/WHEN/THEN scenario.", "`specs lint` fails on any gap."]),
    ("2. Human PR Gate", ["Spec changes go through a pull request that needs BA approval.", "No planning until it merges."]),
    ("3. Grounded in the UI", ["Screenshots in `docs/specs/<cap>/assets/`.", "The agent reads fields, rules, and error states from them."]),
]):
    card(s, x, y, w, h, t, paras=ps)
callout(s, L, 5.6, CW, 0.75, "Human in the loop:", "The agent drafts; the Business Analyst decides.", "human")

# ---------- Slide 6 ----------
s = new_slide("""
Stage 2 và 3: Plan và Orchestrate.

Track có giới hạn:
- Mỗi track trong plan.md gắn với một thư mục, ví dụ Area: modern_api hoặc Area: legacy_core.
- Mọi chỉnh sửa ngoài Area đã khai báo đều bị từ chối.
- Mỗi track ánh xạ tới requirement cụ thể, ví dụ Requirements: REQ-AUTH-001.

Cô lập bằng worktree:
- /orchestrate tạo một Git worktree riêng dưới .planning/worktrees/ cho từng track.
- Nhiều developer và nhiều agent có thể build, test song song mà không tranh chấp nhánh.
- Khi PR upstream được merge, các track downstream tự động được mở khoá trong Beads.

Định tuyến model theo độ phức tạp:
- Orchestrator đánh giá độ khó về cấu trúc và thuật toán của từng track, rồi đề xuất mức reasoning (low, medium, high) và model phù hợp.
- Tránh vòng lặp retry khi model quá yếu, và tránh lãng phí model mạnh cho code boilerplate.
""")
header(s, "Stages 2 & 3: Plan & Orchestrate", "Bounded Tracks, Isolated Worktrees, Routed Models",
       "Each track has a fixed scope, its own workspace, and a model sized to the work.")
for (x, y, w, h), (t, its) in zip(cols(3, 2.1, 4.3), [
    ("Bounded Tracks", ["Each track owns one `Area:` directory.", "Edits outside that area are rejected.", "Each track maps to requirement IDs."]),
    ("Worktree Isolation", ["One Git worktree per track.", "Parallel work without branch contention.", "Downstream tracks unlock as upstream merges."]),
    ("Model Routing", ["Track complexity sets the reasoning tier.", "No under-powered retry loops.", "No premium models on boilerplate."]),
]):
    card(s, x, y, w, h, t, items=its)

# ---------- Slide 7 ----------
s = new_slide("""
Stage 4: Execute. Đi theo chuỗi trên sơ đồ: requirement REQ-AUTH-003 trong living spec → track trong plan.md → test thất bại được gắn tag # REQ-AUTH-003 → gate specs trace.

- Developer nhận track bằng lệnh gin-workflow team claim <bead>.
- Không chấp nhận code triển khai nếu chưa có unit test và integration test thất bại trước, gắn đúng mã requirement.
- Trước khi hoàn tất, gin-workflow specs trace đối chiếu từng scenario của requirement với mã nguồn test. Scenario nào thiếu test sẽ chặn gate ship.
""")
header(s, "Stage 4: Execute", "Code Traced Back to the Spec",
       "Every requirement is tagged from spec to test and audited before ship.")
k = diagram_scale(940)
oy = 2.15
for x, w, fill, stroke, wt, t1, c1, sub, c2 in (
    (20, 160, C["bg"], C["ink"], 1.2, "REQ-AUTH-003", C["agent"], "Living Spec Contract", C["ink"]),
    (260, 160, C["bg"], C["ink"], 1.2, "plan.md Track", C["agent"], "Assigned to Engineer", C["ink"]),
    (500, 180, C["band_agents"], C["agent"], 1.2, "Test: # REQ-AUTH-003", C["agent"], "TDD Failing Test First", C["ink"]),
    (760, 160, C["success_soft"], C["success"], 1.5, "specs trace Gate", C["success"], "100% Scenario Audit", C["success"]),
):
    dbox(s, L, oy, k, x, 15, w, 60, fill, stroke, wt,
         [(t1, 12, c1, MONO, True), (sub, 11, c2, BODY, False)], 1.15)
for x1, x2 in ((180, 260), (420, 500), (680, 760)):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(L + x1 * k), Inches(oy + 45 * k),
                                Inches(L + x2 * k), Inches(oy + 45 * k))
    ln.line.color.rgb = C["ink"]
    ln.line.width = Pt(1.5)
    etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
a, b = cols(2, 3.55, 2.2)
card(s, *a, "Test First", paras=["A failing test tagged with the requirement ID comes before any implementation."])
card(s, *b, "`specs trace` Gate", paras=["Every scenario must map to a test. Gaps block ship."])

# ---------- Slide 8 ----------
s = new_slide("""
Review độc lập:
- Một review agent tách biệt kiểm tra diff của track dựa trên quy tắc kiến trúc, độ phủ test và tiêu chí bảo mật.
- Findings được lưu trong .planning/reviews/<bead>/findings.jsonl, kèm số dòng và mức rủi ro.
- Số vòng review có giới hạn, tránh tranh luận vòng vo.

Engineer không chấp nhận finding của AI một cách thụ động. Mỗi finding phải được kiểm chứng kỹ thuật và có một trong ba quyết định:
- Fix: sửa code và baseline test, rồi kiểm chứng lại.
- Dispute: đưa ra lập luận kỹ thuật cho thấy code hiện tại đúng như thiết kế.
- Defer: ghi lại issue chính thức hoặc waiver kiến trúc có người ký.

Quy tắc: track không được đóng khi ledger còn finding chưa xử lý.
""")
header(s, "Verification", "Independent Review, Recorded in a Ledger", "Every finding gets a documented decision.")
a, b = cols(2, 2.1, 3.25)
card(s, *a, "Independent Review", items=[
    "A separate reviewer agent checks the diff.", "Findings are logged with line and risk.", "Review cycles are capped."])
card(s, *b, "Every Finding Resolved", items=[
    "**Fix:** correct and re-verify.", "**Dispute:** show why the code is right.", "**Defer:** record an issue or waiver."])
callout(s, L, 5.6, CW, 0.75, "Rule:", "No track closes with open findings.")

# ---------- Evidence ----------
s = new_slide("""
Nguyên tắc: không agent nào tự chứng nhận công việc của mình. Mỗi tuyên bố cần bằng chứng quan sát được.

Đọc bảng theo từng hàng:
- "Tests pass" cần output test với 0 failure ở lần chạy mới; một lần chạy trước đó là chưa đủ.
- "Bug fixed" cần test tái hiện đúng triệu chứng ban đầu, đỏ trước rồi xanh; chỉ thay đổi code là chưa đủ.
- "Agent completed" cần diff đã được kiểm tra; agent tự nói là chưa đủ.
- "Requirements met" cần checklist requirement đã được đối chiếu; test pass là chưa đủ.
- "Waiting for merge" cần nhánh đã push và link PR thật; code chưa commit là chưa đủ.

Quy tắc bàn giao: mỗi báo cáo phải nêu rõ đã thay đổi gì, đã kiểm chứng gì, đã miễn trừ gì kèm lý do, trạng thái bead, rủi ro còn lại và lệnh tiếp theo.
""")
header(s, "Verification", "Evidence Over Self-Assessment", "No agent certifies its own work; every claim needs proof.")
table(s, [
    ("Claim", "Requires", "Not enough"),
    ("Tests pass", "Test output with 0 failures", "An earlier run"),
    ("Bug fixed", "The symptom's test goes red, then green", "Code changed"),
    ("Agent completed", "The diff checked", "The agent says so"),
    ("Requirements met", "The checklist checked", "Tests passing"),
    ("Waiting for merge", "A pushed branch and a real PR link", "Uncommitted code"),
], 2.1, (CW * 0.26, CW * 0.44, CW * 0.30), 0.5, 0.52)
callout(s, L, 5.7, CW, 0.75, "Handoff rule:",
        "Every report states what changed, what was validated, what was waived, and what risk remains.")

# ---------- Slide 9 ----------
s = new_slide("""
Cái bẫy khi để AI sinh test: công cụ không bị ràng buộc sẽ sinh hàng trăm test case vụn vặt, giá trị thấp. QA bị ngập trong việc phân loại, làm chậm cả quá trình giao hàng.

Cách gin-qa xử lý:
- qa/guidelines.md đặt chính sách: chỉ hành trình người dùng quan trọng mới là Type: e2e; edge case dùng unit hoặc integration test.
- Mỗi test case lưu hash của requirement nó kiểm tra (REQ-AUTH-001@hash).
- gin-qa cases plan phân loại test case thành:
  - missing: requirement mới, cần thêm test case.
  - stale: requirement đã thay đổi từ lần review trước.
  - obsolete: requirement đã bị xoá.
- Test case không đổi thì giữ nguyên; QA chỉ review phần chênh lệch.
""")
header(s, "Quality Engineering", "Focused QA with `gin-qa`",
       "Test what the requirements say, and re-review only what changed.")
a, b = cols(2, 2.1, 4.3)
card(s, *a, "The Test Generation Trap",
     paras=["Unconstrained AI produces hundreds of low-value tests, and QA drowns in triage."])
callout(s, a[0] + 0.22, 4.25, a[2] - 0.44, 1.9, "The Fix:",
        "Generate tests only from formal requirements, at the depth `qa/guidelines.md` sets.")
card(s, *b, "Hash-Pinned Drift Detection",
     paras=["Each case pins its requirement's hash. `gin-qa cases plan` reports:"],
     items=["`missing`: new requirements.", "`stale`: changed requirements.", "`obsolete`: deleted requirements."])

# ---------- Slide 10 ----------
s = new_slide("""
Ba đòn bẩy để tối ưu chi phí token:

1. Shift left: dùng token ở giai đoạn discuss để làm rõ và xác nhận spec mang lại lợi ích cao nhất. Trong hệ thống brownfield, phát hiện thiếu logic muộn sẽ lan sang nhiều module legacy và đốt token vào các vòng debug và review lặp lại.

2. Định tuyến theo độ phức tạp: model quá yếu sẽ thất bại khi tách các phần phức tạp và rơi vào vòng retry; model quá mạnh thì lãng phí ngân sách cho code boilerplate.

3. Audit bằng CLI chạy local: specs lint, specs trace, specs hash và các gate check chạy bằng Python CLI trên máy, không tốn token, chỉ trả về trạng thái gọn thay vì đọc code thô. Việc bàn giao gọn giữa các stage và worktree tách biệt giúp context hội thoại luôn nhỏ.

Tỷ lệ minh hoạ (không phải số đo thực tế): khoảng 1.000 token đầu tư vào làm rõ spec có thể tránh được hơn 50.000 token rework và retry về sau. Nên thay bằng số đo thật từ /report khi có.
""")
header(s, "Token Economics & Governance", "Token Economics: Three Levers",
       "Spend early on clarity, route by complexity, and audit locally.")
for (x, y, w, h), (t, b) in zip(cols(3, 2.1, 2.9), [
    ("1. Shift Left", "Clarifying specs costs far less than late rework rippling through legacy modules."),
    ("2. Complexity Routing", "The right model tier per track means fewer retries and less overspend."),
    ("3. Local CLI Audits", "Lint, trace, hash, and gate checks run locally at zero token cost."),
]):
    card(s, x, y, w, h, t, paras=[b])
callout(s, L, 5.3, CW, 0.85, "Principle:",
        "A small token investment up front prevents large downstream rework and retry costs.", "success")

# ---------- Slide 11 ----------
s = new_slide("""
Stage 5: Ship. Gate ship gồm:
1. Kiểm tra trước khi ship: chạy toàn bộ test, lint, typecheck và specs trace.
2. specs archive gộp requirement từ spec-delta.md vào docs/specs/<cap>/spec.md.
3. Thư mục change proposal được chuyển vào docs/changes/archive/ cùng toàn bộ lịch sử audit.
4. PR tích hợp cuối cùng cần QE duyệt (team.approvals.ship: [qe]).

Sức khoẻ dài hạn:
- docs/specs/ luôn là nguồn sự thật duy nhất cho toàn bộ ứng dụng.
- Khi module mới được ship, các đường code legacy tương ứng được tắt và dọn dẹp.
- Requirement thay đổi cùng với code nên không phát sinh nợ tài liệu; bộ regression test giữ business logic lâu dài.
""")
header(s, "Stage 5: Ship", "Ship Safely, Keep Specs Alive", "Release and spec update happen in one step.")
a, b = cols(2, 2.1, 4.3)
card(s, *a, "The Ship Gate", kind="num", items=[
    "Full tests, lint, typecheck, and `specs trace`.",
    "`specs archive` merges the delta into `spec.md`.",
    "The change folder moves to the archive.",
    "Integration PR needs QE approval."])
card(s, *b, "Long-Term Health", items=[
    "`docs/specs/` stays the source of truth.",
    "Legacy paths retire as modules ship.",
    "Regression tests preserve business logic."])

# ---------- Slide 12 ----------
s = new_slide("""
Tóm tắt bốn kết quả chính:
- Verifiable Specs: ngữ pháp EARS với scenario GIVEN/WHEN/THEN bắt buộc, được BA duyệt qua pull request.
- Parallel Start: phân tích phụ thuộc để chọn epic không bị chặn; sprint đầu chạy song song thiết lập môi trường và xác nhận spec.
- Code Parity: tag requirement hai chiều và audit specs trace bảo đảm code khớp với spec.
- Focused QA: phát hiện drift dựa trên hash giúp QA chỉ tập trung vào phần thay đổi và các hành trình có giá trị cao.

Đề xuất: chạy pilot 2 tuần trên một capability có phạm vi rõ ràng (ví dụ Authentication hoặc Order Intake) để đo velocity cơ sở và chứng minh khả năng giao hàng đạt chuẩn production.
""")
text(s, L, 1.3, CW, 0.3, "SUMMARY & ACTION PLAN", 12, C["agent"], MONO, True)
text(s, L, 1.6, CW, 0.65, "The Modernization Governance Blueprint", 30, C["ink"], DISPLAY, True)
text(s, L, 2.3, CW, 0.4, "Controlled, verifiable delivery for engineering teams and leadership.", 18, C["muted"], DISPLAY)
for (x, y, w, h), (t, b) in zip(cols(4, 3.0, 2.0), [
    ("1. Verifiable Specs", "Strict grammar, approved by BA pull request."),
    ("2. Parallel Start", "Unblocked epics first; setup and specs in parallel."),
    ("3. Code Parity", "Requirement tags and `specs trace` audits."),
    ("4. Focused QA", "Drift detection keeps tests on what changed."),
]):
    card(s, x, y, w, h, t, paras=[b])
callout(s, L, 5.3, CW, 0.8, "Next Step:",
        "A 2-week pilot on one bounded capability, such as Authentication or Order Intake.", "success")

prs.save(DECK_PPTX)
print("slides:", len(prs.slides))
